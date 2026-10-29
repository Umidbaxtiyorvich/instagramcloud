"""Excel fayllar yaratish."""
import io

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HEAD_FILL = PatternFill("solid", fgColor="32348F")
HEAD_FONT = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
THIN = Side(style="thin", color="BBBBBB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NUM = '#,##0'
NUM3 = '#,##0.###'


def _table(ws, start_row: int, headers: list[str], rows: list[list], widths: list[int] | None = None):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=c, value=h)
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r, row in enumerate(rows, start_row + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BORDER
            if isinstance(v, float) and v != int(v):
                cell.number_format = NUM3
            elif isinstance(v, (int, float)):
                cell.number_format = NUM
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return start_row + len(rows) + 1


def _save(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def simple(title: str, headers: list[str], rows: list[list], widths=None, footer: list | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=13)
    end = _table(ws, 3, headers, rows, widths)
    if footer:
        for c, v in enumerate(footer, 1):
            cell = ws.cell(row=end, column=c, value=v)
            cell.font, cell.border = BOLD, BORDER
            if isinstance(v, (int, float)):
                cell.number_format = NUM
    ws.freeze_panes = "A4"
    return _save(wb)


def act_sverka(org_name: str, supplier_name: str, frm, to, act: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Akt sverka"
    ws.merge_cells("A1:F1")
    ws["A1"] = "SOLISHTIRMA DALOLATNOMA (AKT SVERKA)"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A2:F2")
    ws["A2"] = f"{org_name}  va  {supplier_name}"
    ws["A2"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A3:F3")
    ws["A3"] = f"Davr: {frm:%d.%m.%Y} — {to:%d.%m.%Y}"
    ws["A3"].alignment = Alignment(horizontal="center")

    ws["A5"], ws["E5"] = "Boshlang'ich qoldiq:", act["opening"]
    ws["A5"].font = ws["E5"].font = BOLD
    ws["E5"].number_format = NUM

    rows = [[i, r["date"].strftime("%d.%m.%Y %H:%M"), r["document"], r["debit"], r["credit"], r["balance"]]
            for i, r in enumerate(act["rows"], 1)]
    end = _table(ws, 7, ["№", "Sana", "Hujjat", "Debet (to'lov)", "Kredit (tovar)", "Qoldiq"], rows,
                 [6, 18, 48, 18, 18, 18])
    for c, v in enumerate(["", "", "Davr aylanmasi:", act["debit"], act["credit"], ""], 1):
        cell = ws.cell(row=end, column=c, value=v)
        cell.font, cell.border = BOLD, BORDER
        if isinstance(v, (int, float)):
            cell.number_format = NUM
    ws.cell(row=end + 2, column=1, value="Oxirgi qoldiq:").font = BOLD
    cell = ws.cell(row=end + 2, column=5, value=act["closing"])
    cell.font, cell.number_format = BOLD, NUM
    closing = act["closing"]
    if closing < 0:
        txt = f"{org_name} ning {supplier_name} oldidagi qarzi: {abs(round(closing)):,}".replace(",", " ") + " so'm"
    elif closing > 0:
        txt = f"{supplier_name} ning {org_name} oldidagi qarzi: {round(closing):,}".replace(",", " ") + " so'm"
    else:
        txt = "O'zaro qarzdorlik yo'q"
    ws.cell(row=end + 3, column=1, value=txt).font = BOLD
    ws.cell(row=end + 6, column=1, value=f"{org_name}: ____________________")
    ws.cell(row=end + 6, column=4, value=f"{supplier_name}: ____________________")
    return _save(wb)
